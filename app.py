from datetime import datetime
import io
import os
import random
from flask import Flask, jsonify, render_template, request, session, send_file
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import check_password_hash, generate_password_hash

# Initialize app
app = Flask(__name__)

# Security: Pull secret key from environment variable with a safe local dev fallback
app.secret_key = os.environ.get('SECRET_KEY', 'barangay_binubusan_super_secret_key_2026')

# Security: Rate limiter to block brute-force attacks on login
limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["500 per day", "100 per hour"],
    storage_uri="memory://"
)

# Database path setup
basedir = os.path.abspath(os.path.dirname(__file__))
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(basedir, 'barangay.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)


# --- DATABASE MODELS ---
class Appointment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reference_code = db.Column(db.String(20), unique=True, nullable=False)
    full_name = db.Column(db.String(100), nullable=False)
    document_type = db.Column(db.String(100), nullable=False)
    appointment_date = db.Column(db.String(20), nullable=False)
    time_slot = db.Column(db.String(20), nullable=False)
    status = db.Column(db.String(20), default='Pending')
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    def to_dict(self):
        return {
            'id': self.id,
            'reference_code': self.reference_code,
            'full_name': self.full_name,
            'document_type': self.document_type,
            'appointment_date': self.appointment_date,
            'time_slot': self.time_slot,
            'status': self.status,
        }


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(200), nullable=False)
    role = db.Column(db.String(20), default='staff')

    def to_dict(self):
        return {
            'id': self.id,
            'username': self.username,
            'role': self.role
        }


class AuditLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), nullable=False)
    action = db.Column(db.String(255), nullable=False)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)

    def to_dict(self):
        return {
            'id': self.id,
            'username': self.username,
            'action': self.action,
            'timestamp': self.timestamp.strftime('%Y-%m-%d %H:%M:%S')
        }


with app.app_context():
    db.create_all()
    if not User.query.filter_by(role='super_admin').first():
        hashed_password = generate_password_hash('admin123')
        super_admin = User(
            username='superadmin',
            password_hash=hashed_password,
            role='super_admin'
        )
        db.session.add(super_admin)
        db.session.commit()


def log_action(username, action_text):
    new_log = AuditLog(username=username, action=action_text)
    db.session.add(new_log)
    db.session.commit()


# --- ROUTES ---

@app.route('/')
def home():
    return render_template('index.html')


@app.route('/admin')
def admin_dashboard():
    if not session.get('admin_logged_in'):
        return render_template('admin_login.html')
    return render_template('admin.html')


@app.route('/api/admin-login', methods=['POST'])
@limiter.limit("5 per minute")  # Rate limit login endpoint against brute force
def admin_login():
    data = request.json
    username = data.get('username')
    password = data.get('password')

    user = User.query.filter_by(username=username).first()

    if user and check_password_hash(user.password_hash, password):
        session['admin_logged_in'] = True
        session['username'] = user.username
        session['role'] = user.role
        log_action(user.username, f"Logged into system as {user.role.upper()}")
        return jsonify({'success': True, 'role': user.role})

    return jsonify({'error': 'Invalid username or password.'}), 401


@app.route('/api/admin/session', methods=['GET'])
def get_session():
    if not session.get('admin_logged_in'):
        return jsonify({'logged_in': False}), 401
    return jsonify({
        'logged_in': True,
        'username': session.get('username'),
        'role': session.get('role'),
    })


@app.route('/api/admin/users', methods=['GET'])
def get_admin_users():
    if session.get('role') != 'super_admin':
        return jsonify({'error': 'Unauthorized'}), 403
    users = User.query.all()
    return jsonify([u.to_dict() for u in users])


@app.route('/api/admin/audit-logs', methods=['GET'])
def get_audit_logs():
    if session.get('role') != 'super_admin':
        return jsonify({'error': 'Unauthorized'}), 403
    logs = AuditLog.query.order_by(AuditLog.id.desc()).limit(50).all()
    return jsonify([l.to_dict() for l in logs])


@app.route('/api/admin/create-staff', methods=['POST'])
def create_staff():
    if session.get('role') != 'super_admin':
        return jsonify({'error': 'Unauthorized. Super Admin access required.'}), 403

    data = request.json
    username = data.get('username')
    password = data.get('password')
    role = data.get('role', 'staff')

    if not username or not password:
        return jsonify({'error': 'Username and password are required.'}), 400

    if User.query.filter_by(username=username).first():
        return jsonify({'error': 'Username already exists.'}), 400

    hashed_password = generate_password_hash(password)
    new_user = User(username=username, password_hash=hashed_password, role=role)
    db.session.add(new_user)
    db.session.commit()
    log_action(session.get('username'), f"Created staff account: {username}")

    return jsonify({'success': True, 'message': f'Staff account {username} created successfully!'}), 201


@app.route('/api/admin/users/<int:user_id>', methods=['DELETE'])
def delete_staff(user_id):
    if session.get('role') != 'super_admin':
        return jsonify({'error': 'Unauthorized. Super Admin access required.'}), 403

    user_to_delete = User.query.get_or_404(user_id)
    if user_to_delete.username == session.get('username'):
        return jsonify({'error': 'You cannot delete your own active super admin account.'}), 400

    deleted_username = user_to_delete.username
    db.session.delete(user_to_delete)
    db.session.commit()
    log_action(session.get('username'), f"Removed staff account: {deleted_username}")

    return jsonify({'success': True, 'message': 'Staff account removed successfully!'})


@app.route('/api/admin/change-password', methods=['PUT'])
def change_super_admin_password():
    if session.get('role') != 'super_admin':
        return jsonify({'error': 'Unauthorized. Super Admin access required.'}), 403
    
    data = request.json
    new_password = data.get('new_password')
    
    if not new_password or len(new_password) < 6:
        return jsonify({'error': 'Password must be at least 6 characters long.'}), 400
    
    current_username = session.get('username')
    admin_user = User.query.filter_by(username=current_username).first()
    
    if not admin_user:
        return jsonify({'error': 'User not found.'}), 404
    
    admin_user.password_hash = generate_password_hash(new_password)
    db.session.commit()
    log_action(current_username, "Updated Super Admin password")
    
    return jsonify({'success': True, 'message': 'Super admin password updated successfully!'})


@app.route('/api/admin-logout', methods=['POST'])
def admin_logout():
    if session.get('username'):
        log_action(session.get('username'), "Logged out of the system")
    session.clear()
    return jsonify({'success': True})


@app.route('/api/admin/metrics', methods=['GET'])
def admin_metrics():
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
    
    date_filter = request.args.get('date', datetime.utcnow().strftime('%Y-%m-%d'))
    total = Appointment.query.filter_by(appointment_date=date_filter).count()
    pending = Appointment.query.filter_by(appointment_date=date_filter, status='Pending').count()
    completed = Appointment.query.filter_by(appointment_date=date_filter, status='Completed').count()
    
    return jsonify({'date': date_filter, 'total': total, 'pending': pending, 'completed': completed})


@app.route('/api/admin/download-attendance', methods=['GET'])
def download_attendance():
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
        
    date_filter = request.args.get('date', datetime.utcnow().strftime('%Y-%m-%d'))
    appointments = Appointment.query.filter_by(appointment_date=date_filter, status='Completed').all()
    log_action(session.get('username'), f"Downloaded attendance sheet for {date_filter}")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Document Attendance"
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="1F4E78")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append([f"BARANGAY BINUBUSAN - DOCUMENT CLAIM ATTENDANCE LOG"])
    ws.append([f"Date: {date_filter}"])
    ws.append([])
    ws['A1'].font = title_font
    
    headers = ["No.", "Full Name", "Document Type", "Address", "Contact Number", "Signature / Received By"]
    ws.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    row_idx = 5
    for idx, apt in enumerate(appointments, start=1):
        ws.append([idx, apt.full_name, apt.document_type, "Binubusan, Batangas", "N/A", ""])
        for col_num in range(1, 7):
            c = ws.cell(row=row_idx, column=col_num)
            c.font = Font(name="Arial", size=10)
            c.border = thin_border
            if col_num in [1, 4, 5]:
                c.alignment = Alignment(horizontal="center")
        row_idx += 1
        
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    return send_file(
        file_stream, as_attachment=True,
        download_name=f"Attendance_{date_filter}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route('/api/slots-availability', methods=['GET'])
def slots_availability():
    date_str = request.args.get('date')
    if not date_str:
        return jsonify({})
    bookings = Appointment.query.filter_by(appointment_date=date_str).all()
    slot_counts = {}
    for b in bookings:
        slot_counts[b.time_slot] = slot_counts.get(b.time_slot, 0) + 1
    return jsonify(slot_counts)


@app.route('/api/appointments', methods=['POST'])
def create_appointment():
    data = request.json
    full_name = data.get('full_name')
    document_type = data.get('document_type')
    appointment_date = data.get('appointment_date')
    time_slot = data.get('time_slot')
    captcha_answer = data.get('captcha_answer')

    if str(captcha_answer).strip() != '10':
        return jsonify({'error': 'Robot verification failed!'}), 400

    if time_slot == '12:00 PM - 01:00 PM (Lunch Break)':
        return jsonify({'error': 'Appointments are closed during lunch break.'}), 400

    if not full_name or not document_type or not appointment_date or not time_slot:
        return jsonify({'error': 'All fields are required!'}), 400

    current_count = Appointment.query.filter_by(
        appointment_date=appointment_date, time_slot=time_slot
    ).count()
    if current_count >= 15:
        return jsonify({'error': 'Time slot fully booked (Max 15 people).'}), 400

    ref_code = f'BRGY-{random.randint(1000, 9999)}'
    new_apt = Appointment(
        reference_code=ref_code, full_name=full_name,
        document_type=document_type, appointment_date=appointment_date,
        time_slot=time_slot,
    )
    db.session.add(new_apt)
    db.session.commit()

    return jsonify({'message': 'Appointment booked successfully!', 'reference_code': ref_code}), 201


@app.route('/api/appointments', methods=['GET'])
def get_appointments():
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
    appointments = Appointment.query.order_by(Appointment.id.desc()).all()
    return jsonify([apt.to_dict() for apt in appointments])


@app.route('/api/appointments/track/<ref_code>', methods=['GET'])
def track_appointment(ref_code):
    apt = Appointment.query.filter_by(reference_code=ref_code.strip()).first()
    if not apt:
        return jsonify({'error': 'Reference code not found.'}), 404
    return jsonify(apt.to_dict())


@app.route('/api/appointments/<int:apt_id>/status', methods=['PATCH'])
def update_status(apt_id):
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.json
    apt = Appointment.query.get_or_404(apt_id)
    old_status = apt.status
    new_status = data.get('status')
    apt.status = new_status
    db.session.commit()
    log_action(session.get('username'), f"Updated appointment {apt.reference_code} status from {old_status} to {new_status}")
    return jsonify({'message': 'Status updated successfully!'})


@app.route('/api/appointments/<int:apt_id>', methods=['DELETE'])
def delete_appointment(apt_id):
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
    apt = Appointment.query.get_or_404(apt_id)
    ref_code = apt.reference_code
    db.session.delete(apt)
    db.session.commit()
    log_action(session.get('username'), f"Cleared appointment record {ref_code}")
    return jsonify({'message': 'Record cleared successfully!'})


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)